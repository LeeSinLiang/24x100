"""Money inputs: Ward 5 comparable sales and HUD FY2026 income limits."""
from __future__ import annotations

import io
import json
import re
import statistics
from pathlib import Path
from typing import Any

from . import PIPELINE_VERSION
from . import sources as S
from .fetch import REPO, RawCache

OUT_DIR = REPO / "data" / "money"
MUNICODE = "105"          # Pittsburgh Ward 5
SINCE = "2023-01-01"
MIN_PRICE = 10000
VALID = "VALID SALE"
USES_1_2 = ("ROWHOUSE", "SINGLE FAMILY", "TOWNHOUSE", "TWO FAMILY")
QUANTILE_METHOD = (
    "statistics.quantiles(prices, n=4, method='inclusive'): linear interpolation between order "
    "statistics, Hyndman-Fan type 7 (numpy/Excel QUARTILE.INC default)"
)
HUD_AREA = "METRO38300M38300"
HUD_COUNTY = "Allegheny"


# ---------------------------------------------------------------------------------------------
# comparables


def fetch_money(cache: RawCache) -> None:
    keys = S.fetch_sales(cache, MUNICODE)
    recs = _sales_records(cache, keys)
    pins = sorted({r["PARID"] for r in recs if _is_valid_recent(r)})
    S.fetch_assessments(cache, f"sales_{MUNICODE}", pins)
    S.fetch_hud(cache)


def _sales_keys(cache: RawCache) -> list[str]:
    return sorted(k for k in cache.manifest if k.startswith(f"sales_{MUNICODE}/page_"))


def _sales_records(cache: RawCache, keys: list[str]) -> list[dict]:
    recs: list[dict] = []
    total = None
    for k in keys:
        res = cache.read_json(k)["result"]
        total = res["total"]
        recs += res["records"]
    if total is not None and len(recs) != total:
        raise RuntimeError(f"sales pages hold {len(recs)} records but the API total is {total}")
    return recs


def _is_valid_recent(r: dict) -> bool:
    return (
        r.get("SALEDESC") == VALID
        and r.get("PRICE") is not None
        and float(r["PRICE"]) >= MIN_PRICE
        and (r.get("SALEDATE") or "") >= SINCE
    )


def _street_addr(full: str | None) -> str | None:
    """'2125 ROSE ST, PITTSBURGH, PA 15219' -> '2125 Rose St' (property address, no names)."""
    if not full:
        return None
    street = full.split(",")[0]
    street = re.sub(r"\s+", " ", street).strip()
    return " ".join(w.capitalize() if not w[0].isdigit() else w for w in street.split())


def _num(v):
    if v is None or v == "":
        return None
    f = float(v)
    return int(f) if f.is_integer() else f


def build_snapshot(cache: RawCache) -> dict[str, Any]:
    """Filtered, joined records plus the upstream counts: the input to summarize()."""
    sales_keys = _sales_keys(cache)
    recs = _sales_records(cache, sales_keys)
    assess: dict[str, dict] = {}
    for k in sorted(k for k in cache.manifest if k.startswith(f"sales_{MUNICODE}/assessments_")):
        for a in cache.read_json(k)["result"]["records"]:
            assess[a["PARID"]] = a
    records = []
    for r in recs:
        if not _is_valid_recent(r):
            continue
        a = assess.get(r["PARID"])
        records.append({
            "parid": r["PARID"],
            "address": _street_addr(r.get("FULL_ADDRESS")),
            "saledate": r["SALEDATE"][:10],
            "price": _num(r["PRICE"]),
            "saledesc": r["SALEDESC"],
            "instrument": r.get("INSTRTYPDESC"),
            "use": a.get("USEDESC") if a else None,
            "yearbuilt": _num(a.get("YEARBLT")) if a else None,
            "sqft": _num(a.get("FINISHEDLIVINGAREA")) if a else None,
            "assessment_found": a is not None,
        })
    records.sort(key=lambda x: (x["saledate"], x["parid"], x["price"]))
    first = cache.entry(sales_keys[0])
    return {
        "meta": {
            "source": "WPRDC Allegheny County Real Estate Sales + Property Assessments",
            "sales_resource_id": S.SALES_RESOURCE,
            "assessment_resource_id": S.ASSESSMENT_RESOURCE,
            "pulled": first["fetched_at"],
            "municode": MUNICODE,
            "record_filter": {"SALEDESC": VALID, "PRICE_gte": MIN_PRICE, "SALEDATE_gte": SINCE},
        },
        "upstream": {
            "transfers_all_dates": len(recs),
            "transfers_all_dates_first": min(r["SALEDATE"][:10] for r in recs),
            "transfers_all_dates_last": max(r["SALEDATE"][:10] for r in recs),
            "valid_all_dates": sum(1 for r in recs if r.get("SALEDESC") == VALID),
            "transfers_since": sum(1 for r in recs if (r.get("SALEDATE") or "") >= SINCE),
            "valid_since": sum(1 for r in recs if r.get("SALEDESC") == VALID and (r.get("SALEDATE") or "") >= SINCE),
        },
        "records": records,
    }


def summarize(snapshot: dict[str, Any]) -> dict[str, Any]:
    """Counts, median, IQR and newest comparables from a snapshot. Pure; used by tests."""
    recs = snapshot["records"]
    keep = [r for r in recs if r["use"] in USES_1_2]
    prices = sorted(r["price"] for r in keep)
    if len(prices) < 2:
        raise ValueError("fewer than two comparable sales; cannot compute an IQR")
    q1, med, q3 = statistics.quantiles(prices, n=4, method="inclusive")
    alt = statistics.quantiles(prices, n=4, method="exclusive")
    newest = sorted(
        (r for r in keep if r["yearbuilt"] is not None),
        key=lambda r: (-r["yearbuilt"], _neg_date(r["saledate"]), r["parid"]),
    )[:3]
    up = snapshot["upstream"]

    def money(v):
        return int(v) if float(v).is_integer() else round(v, 2)

    return {
        "counts": {
            "transfers": up["transfers_since"],
            "valid": len(recs),
            "valid_1_2_unit": len(keep),
            "window": [SINCE, up["transfers_all_dates_last"]],
            "valid_any_price": up["valid_since"],
            "transfers_all_dates": up["transfers_all_dates"],
            "valid_all_dates": up["valid_all_dates"],
            "all_dates_range": [up["transfers_all_dates_first"], up["transfers_all_dates_last"]],
            "note": (
                f"Since {SINCE}: {up['transfers_since']} Ward 5 transfers, {up['valid_since']} valid sales "
                f"({len(recs)} at >= ${MIN_PRICE:,}), {len(keep)} of them 1-2 unit homes. The whole sales "
                f"file ({up['transfers_all_dates_first']} to {up['transfers_all_dates_last']}) holds "
                f"{up['transfers_all_dates']} Ward 5 transfers, {up['valid_all_dates']} of them valid sales."
            ),
        },
        "median": money(med),
        "q1": money(q1),
        "q3": money(q3),
        "quantile_method": QUANTILE_METHOD,
        "quantile_alternatives": {
            "exclusive_type6": {"q1": money(alt[0]), "median": money(alt[1]), "q3": money(alt[2])},
        },
        "newest": _newest(newest[0]) if newest else None,
        "newest_built": [_newest(r) for r in newest],
        "newest_rule": "top 3 by assessment YEARBLT (newest first), ties by sale date (newest first)",
        "sales": [
            {"parid": r["parid"], "address": r["address"], "saledate": r["saledate"], "price": r["price"],
             "use": r["use"], "yearbuilt": r["yearbuilt"], "sqft": r["sqft"]}
            for r in keep
        ],
        "excluded_by_use": [
            {"parid": r["parid"], "address": r["address"], "saledate": r["saledate"], "price": r["price"], "use": r["use"]}
            for r in recs if r["use"] not in USES_1_2
        ],
    }


def _newest(r: dict) -> dict:
    return {"addr": r["address"], "parid": r["parid"], "yearbuilt": r["yearbuilt"], "price": r["price"],
            "sqft": r["sqft"], "saledate": r["saledate"], "use": r["use"]}


def _neg_date(d: str) -> str:
    # sort newest first inside a stable ascending sort key
    return "".join(chr(0x7F - ord(c)) if c.isdigit() else c for c in d)


def comps_file(snapshot: dict[str, Any]) -> dict[str, Any]:
    s = summarize(snapshot)
    meta = dict(snapshot["meta"])
    meta.pop("record_filter", None)
    meta.update({
        "filters": {
            "MUNICODE": MUNICODE, "ward": 5, "SALEDESC": VALID, "PRICE_gte": MIN_PRICE,
            "SALEDATE_gte": SINCE, "USEDESC_in": list(USES_1_2),
        },
        "resource_id": S.SALES_RESOURCE,
        "url": f"{S.CKAN_SEARCH}?resource_id={S.SALES_RESOURCE}",
        "pipeline_version": PIPELINE_VERSION,
        "privacy": "property addresses only; the sales file carries no party names and none are added",
    })
    out = {"meta": meta}
    out.update(s)
    return out


# ---------------------------------------------------------------------------------------------
# HUD


def hud_file(cache: RawCache) -> dict[str, Any]:
    import openpyxl

    e = cache.entry("hud_il_fy2026")
    wb = openpyxl.load_workbook(io.BytesIO(cache.read("hud_il_fy2026")), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h is not None else "" for h in next(rows)]
    lower = [h.lower() for h in hdr]
    hits = []
    for r in rows:
        d = dict(zip(lower, r))
        county = str(d.get("county_name") or d.get("countyname") or "")
        if d.get("hud_area_code") == HUD_AREA and HUD_COUNTY.lower() in county.lower():
            hits.append(d)
    if len(hits) != 1:
        raise RuntimeError(f"HUD file: expected one {HUD_AREA} / {HUD_COUNTY} row, found {len(hits)}")
    d = hits[0]
    median_key = next(k for k in lower if k.startswith("median"))
    out = {
        "meta": {
            "source": "HUD FY2026 Income Limits (Section 8), huduser.gov",
            "url": e["url"],
            "file_sha256": e["sha256"],
            "pulled": e["fetched_at"],
            "sheet": wb.sheetnames[0],
            "pipeline_version": PIPELINE_VERSION,
        },
        "hud_area_code": d["hud_area_code"],
        "hud_area_name": d.get("hud_area_name"),
        "county": d.get("county_name") or d.get("countyname"),
        "state": d.get("state_alpha") or d.get("stusps"),
        "median_family_income": int(d[median_key]),
        "median_field": hdr[lower.index(median_key)],
    }
    for lvl in ("l50", "l80"):
        for n in range(1, 9):
            out[f"{lvl}_{n}"] = int(d[f"{lvl}_{n}"])
    return out


def dumps(obj: Any) -> str:
    return json.dumps(obj, sort_keys=True, indent=1, ensure_ascii=False) + "\n"


def write(obj: Any, name: str, out_dir: Path = OUT_DIR) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    p = out_dir / name
    p.write_text(dumps(obj))
    return p
